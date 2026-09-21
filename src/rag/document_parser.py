"""
Document Parser
Parses multiple file formats (PDF, Word, PPT, Excel, TXT, Markdown, HTML) into plain text.

Supported formats:
- .pdf  → pypdf
- .docx → python-docx
- .doc  → (reads as binary, attempts basic extraction)
- .pptx → python-pptx
- .xlsx → openpyxl
- .txt, .md → direct read
- .html → BeautifulSoup
"""
import os
import io
from typing import BinaryIO


class DocumentParser:
    """Parses uploaded documents into plain text for RAG indexing."""

    SUPPORTED_EXTENSIONS = {
        ".pdf", ".docx", ".doc", ".pptx", ".xlsx",
        ".txt", ".md", ".html", ".htm",
    }

    def parse(self, filename: str, file: BinaryIO) -> str:
        """
        Parse a file and return plain text content.

        Args:
            filename: Original filename (used to detect format by extension)
            file: File-like object (binary)

        Returns:
            Extracted plain text
        """
        ext = os.path.splitext(filename)[1].lower()

        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format: {ext}. "
                f"Supported: {', '.join(sorted(self.SUPPORTED_EXTENSIONS))}"
            )

        parser_method = {
            ".pdf": self._parse_pdf,
            ".docx": self._parse_docx,
            ".doc": self._parse_doc,
            ".pptx": self._parse_pptx,
            ".xlsx": self._parse_xlsx,
            ".txt": self._parse_text,
            ".md": self._parse_text,
            ".html": self._parse_html,
            ".htm": self._parse_html,
        }

        return parser_method[ext](file)

    def _parse_pdf(self, file: BinaryIO) -> str:
        """Extract text from PDF."""
        from pypdf import PdfReader

        reader = PdfReader(file)
        texts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                texts.append(text.strip())
        return "\n\n".join(texts)

    def _parse_docx(self, file: BinaryIO) -> str:
        """Extract text from Word .docx."""
        from docx import Document

        doc = Document(file)
        texts = []
        for para in doc.paragraphs:
            if para.text.strip():
                texts.append(para.text.strip())

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    texts.append(row_text)

        return "\n".join(texts)

    def _parse_doc(self, file: BinaryIO) -> str:
        """Attempt to extract text from legacy .doc format."""
        # Legacy .doc is binary; attempt basic text extraction
        raw = file.read()
        try:
            # Try decoding as text, stripping non-printable chars
            text = raw.decode("utf-8", errors="ignore")
            # Filter to printable characters
            import re
            text = re.sub(r"[^\x20-\x7E\n\r\t]", " ", text)
            text = re.sub(r" {3,}", "  ", text)
            return text.strip() if text.strip() else "[Could not extract text from .doc file]"
        except Exception:
            return "[Could not extract text from .doc file]"

    def _parse_pptx(self, file: BinaryIO) -> str:
        """Extract text from PowerPoint .pptx."""
        from pptx import Presentation

        prs = Presentation(file)
        texts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            slide_texts.append(row_text)
            if slide_texts:
                texts.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_texts))

        return "\n\n".join(texts)

    def _parse_xlsx(self, file: BinaryIO) -> str:
        """Extract text from Excel .xlsx."""
        from openpyxl import load_workbook

        wb = load_workbook(file, read_only=True, data_only=True)
        texts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            texts.append(f"--- Sheet: {sheet} ---")
            for row in ws.iter_rows(values_only=True):
                row_text = " | ".join(str(cell).strip() for cell in row if cell is not None)
                if row_text:
                    texts.append(row_text)
        wb.close()
        return "\n".join(texts)

    def _parse_text(self, file: BinaryIO) -> str:
        """Read plain text or Markdown."""
        content = file.read()
        # Try UTF-8, fall back to latin-1
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1")

    def _parse_html(self, file: BinaryIO) -> str:
        """Extract text from HTML."""
        from bs4 import BeautifulSoup

        content = file.read()
        html = content.decode("utf-8", errors="ignore")
        soup = BeautifulSoup(html, "html.parser")

        # Remove script and style elements
        for tag in soup(["script", "style"]):
            tag.decompose()

        return soup.get_text(separator="\n", strip=True)


# Singleton
document_parser = DocumentParser()
